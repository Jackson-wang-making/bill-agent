"""Word：提取正文与表格（.docx=python-docx；.doc 走 LibreOffice 转换，无则降级）。"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path


def _parse_docx(path: str) -> dict:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables: list[list[list[str]]] = []
    for t in doc.tables:
        tables.append([[cell.text for cell in row.cells] for row in t.rows])

    return {
        "type": "word",
        "source_file": str(path),
        "raw_text": "\n".join(paragraphs),
        "tables": tables,
        "meta": {},
    }


def _convert_doc_to_docx(path: str) -> str:
    """LibreOffice headless 把 .doc 转 .docx，返回转换后路径；失败抛带真实原因的异常。"""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("旧版 .doc 需 LibreOffice headless 转换，当前环境未安装")
    outdir = Path(tempfile.mkdtemp(prefix="bill_doc_"))
    try:
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", str(outdir), str(path)],
                check=True, timeout=120, capture_output=True,
            )
        except Exception as e:  # noqa: BLE001 - 崩溃/超时与「未安装」是不同的真实原因
            raise RuntimeError(f"旧版 .doc 转换失败（LibreOffice）: {e}") from e
        converted = list(outdir.glob("*.docx"))
        if not converted:
            raise RuntimeError("旧版 .doc 转换失败（LibreOffice 未产出 .docx）")
    except Exception:
        # 转换失败也必须清理：否则每次失败都在 %TEMP% 残留一个 bill_doc_* 目录
        shutil.rmtree(outdir, ignore_errors=True)
        raise
    return str(converted[0])


def parse_word(path: str) -> dict:
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return _parse_docx(path)

    # .doc 旧二进制：python-docx 不支持，转 docx 后解析；转换失败转人工并写真实原因
    try:
        converted = _convert_doc_to_docx(path)
    except Exception as e:  # noqa: BLE001 - 真实原因必须进审计轨迹
        return {
            "type": "word",
            "source_file": str(path),
            "raw_text": "",
            "tables": [],
            "meta": {"error": f"{e}，转入 PENDING_REVIEW"},
        }
    try:
        result = _parse_docx(converted)
    finally:
        # 转换产物只在本次解析内需要：bill_doc_ 临时目录不得残留
        parent = Path(converted).parent
        if parent.name.startswith("bill_doc_"):
            shutil.rmtree(parent, ignore_errors=True)
    result["meta"]["converted_from_doc"] = True
    # source_file 必须还原为原始 .doc 路径：转换落盘的 bill_doc_ 临时目录
    # 不得进入审计轨迹（original_text 头部）或 LLM prompt
    result["source_file"] = str(path)
    return result
