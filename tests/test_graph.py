"""图编译 + 无 key 时的 harness 降级（转 PENDING_REVIEW）。"""
from bill_agent.graph.builder import build_graph
from bill_agent.main import run_bill_agent


def test_trace_node_returns_result_and_logs_summary(caplog):
    """节点追踪透传返回值，并记录耗时与输出键摘要（不序列化 state 值）。"""
    import logging

    from bill_agent.harness.logger import trace_node

    with caplog.at_level(logging.INFO):
        out = trace_node("probe", lambda state: {"output_records": [1]})(
            {"parsed_documents": [{"raw_text": "secret"}]}
        )

    assert out == {"output_records": [1]}
    assert "elapsed=" in caplog.text
    assert "out=[output_records]" in caplog.text


def test_second_confirm_passes_full_extraction_result(monkeypatch):
    """语义复核必须拿到 document_info/交易方，不能只传 charges 导致误报缺字段。"""
    import bill_agent.graph.nodes.validate_extract as ve

    captured = {}

    def fake_check(result, doc_text):
        captured.update(result)
        return None

    monkeypatch.setattr(ve, "semantic_check", fake_check)
    out = ve.second_confirm_node({
        "document_info": {"document_type": "PAI-CMA-NOTICE"},
        "transaction_party": {"issuer": "A", "recipient": "B"},
        "extracted_charges": [{"charge_item": "On Carriage", "amount": 245.0,
                               "currency": "USD"}],
        "subtotal": None, "tax": None, "total": None,
        "parsed_documents": [], "iteration": 1,
    })

    assert captured["document_info"]["document_type"] == "PAI-CMA-NOTICE"
    assert captured["transaction_party"]["issuer"] == "A"
    assert out["second_confirm_status"] == "pass"


def test_pending_review_preserves_extracted_charge():
    """语义复核失败转人工时，已提取的费用必须保留供人工核对，不能清空。"""
    from bill_agent.graph.nodes.pending_review import pending_review_node

    out = pending_review_node({
        "validation_status": "fail",
        "extraction_confidence": {"overall": 0.7},
        "extracted_charges": [{
            "charge_item": "On Carriage Additional - US Em",
            "original_charge_item": "On Carriage Additional - US Em",
            "amount": 245.0, "currency": "USD", "payment_term": None,
            "quantity": None, "unit_price": None,
        }],
        "document_info": {"document_type": "Arrival Notice"},
        "transaction_party": {"issuer": "CMA CGM", "recipient": "NEPTUNE"},
        "parsed_documents": [{
            "type": "pdf_text", "source_file": "x.pdf",
            "raw_text": "On Carriage USD 245.00", "tables": [], "meta": {},
        }],
        "accounting_entity": "NEPTUNE",
    })

    details = out["output_records"][0]["charge_details"]
    assert details["charge_item"] == "On Carriage Additional - US Em"
    assert details["amount"] == 245.0
    assert details["currency"] == "USD"


def test_graph_compiles():
    g = build_graph()
    assert g.compile() is not None


def test_build_checkpointer_default_is_none():
    """默认不启用 checkpointer：单次 CLI 与离线测试行为保持不变。"""
    from bill_agent.main import build_checkpointer

    assert build_checkpointer() is None


def test_build_checkpointer_memory(monkeypatch):
    from types import SimpleNamespace

    import bill_agent.main as main
    from langgraph.checkpoint.memory import InMemorySaver

    monkeypatch.setattr(
        main, "config",
        SimpleNamespace(checkpointer="memory", checkpointer_sqlite_path=""),
    )
    assert isinstance(main.build_checkpointer(), InMemorySaver)


def test_build_checkpointer_sqlite_creates_file(monkeypatch, tmp_path):
    from types import SimpleNamespace

    import bill_agent.main as main
    from langgraph.checkpoint.sqlite import SqliteSaver

    db = tmp_path / "checkpoints.sqlite"
    monkeypatch.setattr(
        main, "config",
        SimpleNamespace(checkpointer="sqlite", checkpointer_sqlite_path=str(db)),
    )

    saver = main.build_checkpointer()
    assert isinstance(saver, SqliteSaver)
    assert db.exists()


def test_run_bill_agent_with_checkpointer(monkeypatch, tmp_path):
    """带 checkpointer 编译图后仍能跑通整条管线，不改变输出。"""
    import sqlite3

    from langgraph.checkpoint.sqlite import SqliteSaver
    from tests.test_e2e_mock import _mock_extract, _result, _write_docx

    _mock_extract(monkeypatch, _result())
    records = run_bill_agent(
        [_write_docx(tmp_path)], [], "A Shipping",
        checkpointer=SqliteSaver(sqlite3.connect(":memory:", check_same_thread=False)),
    )
    assert [r["operation"] for r in records] == ["CREATE"]


def test_unknown_file_routes_to_pending_review(tmp_path):
    """未知类型 → parse 错误 → aggregate 短路 → PENDING_REVIEW（不空跑 LLM）。"""
    p = tmp_path / "a.xyz"
    p.write_text("hello", encoding="utf-8")
    records = run_bill_agent([str(p)])
    assert records
    assert records[0]["operation"] == "PENDING_REVIEW"
    # 错误与文件名必须进审计轨迹（吞掉错误会让配置故障伪装成成功转人工）
    trail = records[0]["audit_trail"]
    assert trail["source_file"] == str(p)
    assert any("无法识别的文件类型" in line for line in trail["ai_reasoning"])


def test_image_parse_failure_carries_real_error(monkeypatch, tmp_path):
    """图片 + OSS 故障：PENDING_REVIEW 记录必须带真实错误与文件名。"""
    import bill_agent.graph.nodes.parse as parse_mod

    def boom(path):
        raise RuntimeError("OSS 权限不足：当前 RAM 用户对桶 rag-data-2026 没有 oss:PutObject 权限")

    monkeypatch.setattr(parse_mod, "_parse_ocr", boom)

    p = tmp_path / "发票.png"
    p.write_bytes(b"\x89PNG\r\n")
    records = run_bill_agent([str(p)])

    assert len(records) == 1
    rec = records[0]
    assert rec["operation"] == "PENDING_REVIEW"
    assert rec["audit_trail"]["source_file"] == str(p)
    joined = "\n".join(rec["audit_trail"]["ai_reasoning"])
    assert "oss:PutObject" in joined and str(p) in joined


def test_partial_parse_failure_keeps_normal_records_and_errors(monkeypatch, tmp_path):
    """部分文件解析失败：正常文件照常出单，失败文件补 PENDING_REVIEW（错误不丢）。"""
    from tests.test_e2e_mock import _mock_extract, _result, _write_docx

    _mock_extract(monkeypatch, _result())

    bad = tmp_path / "bad.xyz"
    bad.write_text("hello", encoding="utf-8")
    good = _write_docx(tmp_path)

    records = run_bill_agent([str(good), str(bad)], [], "A Shipping")

    ops = [r["operation"] for r in records]
    assert "CREATE" in ops
    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert len(pending) == 1
    assert pending[0]["audit_trail"]["source_file"] == str(bad)
    assert any("无法识别的文件类型" in line for line in pending[0]["audit_trail"]["ai_reasoning"])


def test_eml_attachment_failure_surfaces_filename_not_temp_path(monkeypatch, tmp_path):
    """EML 附件解析失败：错误与附件文件名进 PENDING_REVIEW，不吞掉、不泄漏临时目录路径。"""
    import bill_agent.parsers.ocr as ocr_mod
    from tests.test_e2e_mock import _mock_extract, _result

    def boom(path):
        raise RuntimeError("OSS 权限不足：当前 RAM 用户对桶 rag-data-2026 没有 oss:PutObject 权限")

    # EML 解析器在解析附件时才 from .ocr import parse_ocr，patch 模块属性即可命中
    monkeypatch.setattr(ocr_mod, "parse_ocr", boom)
    _mock_extract(monkeypatch, _result())

    boundary = "BOUNDARY123"
    eml = tmp_path / "mail.eml"
    eml.write_bytes(
        (
            f"Content-Type: multipart/mixed; boundary={boundary}\r\n"
            "Subject: PRE-ALERT\r\nFrom: a@b.com\r\nTo: c@d.com\r\n"
            f"\r\n--{boundary}\r\nContent-Type: text/plain\r\n\r\nInvoice attached.\r\n"
            f'--{boundary}\r\nContent-Type: image/png\r\nContent-Disposition: attachment; filename="发票.png"\r\n\r\n'
        ).encode("utf-8")
        + b"\x89PNG\r\n"
        + f"\r\n--{boundary}--\r\n".encode("utf-8")
    )

    records = run_bill_agent([str(eml)], [], "A Shipping")

    ops = [r["operation"] for r in records]
    assert "CREATE" in ops
    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert len(pending) == 1
    trail = pending[0]["audit_trail"]
    assert trail["source_file"] == "发票.png"
    joined = "\n".join(trail["ai_reasoning"])
    assert "发票.png" in joined and "oss:PutObject" in joined
    # 附件临时目录路径不得泄漏进任何记录的审计轨迹
    for r in records:
        t = r["audit_trail"]
        blob = "\n".join([str(t.get("source_file")), str(t.get("original_text")),
                          "\n".join(t.get("ai_reasoning") or [])])
        assert "bill_att_" not in blob


def _assert_no_temp_leak(records):
    """任何记录的审计轨迹（source_file/original_text/reasoning）不得出现附件临时目录路径。"""
    for r in records:
        t = r["audit_trail"]
        blob = "\n".join([str(t.get("source_file")), str(t.get("original_text")),
                          "\n".join(t.get("ai_reasoning") or [])])
        assert "bill_att_" not in blob


def test_eml_unknown_attachment_does_not_leak_temp_path(monkeypatch, tmp_path):
    """未知类型附件（data.bin）：PENDING_REVIEW 带附件文件名，临时目录路径不得泄漏。"""
    from tests.test_e2e_mock import _mock_extract, _result

    _mock_extract(monkeypatch, _result())

    boundary = "BOUNDARY789"
    eml = tmp_path / "mail.eml"
    eml.write_bytes(
        (
            f"Content-Type: multipart/mixed; boundary={boundary}\r\n"
            "Subject: PRE-ALERT\r\nFrom: a@b.com\r\nTo: c@d.com\r\n"
            f"\r\n--{boundary}\r\nContent-Type: text/plain\r\n\r\nInvoice attached.\r\n"
            f'--{boundary}\r\nContent-Type: application/octet-stream\r\nContent-Disposition: attachment; filename="data.bin"\r\n\r\n'
        ).encode("utf-8")
        + b"SOMEBYTES\r\n"
        + f"\r\n--{boundary}--\r\n".encode("utf-8")
    )

    records = run_bill_agent([str(eml)], [], "A Shipping")

    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert len(pending) == 1
    trail = pending[0]["audit_trail"]
    assert trail["source_file"] == "data.bin"
    joined = "\n".join(trail["ai_reasoning"])
    assert "data.bin" in joined and "无法识别文件类型" in joined
    _assert_no_temp_leak(records)


def test_eml_attachment_without_filename_uses_generated_name(monkeypatch, tmp_path):
    """无 filename 的附件失败：用 _safe_filename 生成的 attachment_N，不回落到临时路径。"""
    import bill_agent.parsers.ocr as ocr_mod
    from tests.test_e2e_mock import _mock_extract, _result

    def boom(path):
        raise RuntimeError("OSS 权限不足：当前 RAM 用户对桶 rag-data-2026 没有 oss:PutObject 权限")

    monkeypatch.setattr(ocr_mod, "parse_ocr", boom)
    _mock_extract(monkeypatch, _result())

    boundary = "BOUNDARY101"
    eml = tmp_path / "mail.eml"
    eml.write_bytes(
        (
            f"Content-Type: multipart/mixed; boundary={boundary}\r\n"
            "Subject: PRE-ALERT\r\nFrom: a@b.com\r\nTo: c@d.com\r\n"
            f"\r\n--{boundary}\r\nContent-Type: text/plain\r\n\r\nInvoice attached.\r\n"
            f"--{boundary}\r\nContent-Type: image/png\r\nContent-Disposition: attachment\r\n\r\n"
        ).encode("utf-8")
        + b"\x89PNG\r\n"
        + f"\r\n--{boundary}--\r\n".encode("utf-8")
    )

    records = run_bill_agent([str(eml)], [], "A Shipping")

    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert len(pending) == 1
    trail = pending[0]["audit_trail"]
    # 无文件名附件按 MIME 补扩展名，便于路由与审计辨认
    assert trail["source_file"] == "attachment_0.png"
    joined = "\n".join(trail["ai_reasoning"])
    assert "attachment_0" in joined and "oss:PutObject" in joined
    _assert_no_temp_leak(records)


def test_network_error_with_token_url_is_redacted(monkeypatch, tmp_path):
    """网关异常带 token query（如 ocr_url 配了 api_token）：收口脱敏，token 不进审计轨迹。"""
    import requests

    import bill_agent.graph.nodes.parse as parse_mod

    def boom(path):
        raise requests.HTTPError(
            "502 Server Error for url: https://gateway.pltplt.com/ai/process?api_token=SECRETTOKEN123&x=1"
        )

    monkeypatch.setattr(parse_mod, "_parse_ocr", boom)
    p = tmp_path / "发票.png"
    p.write_bytes(b"\x89PNG\r\n")

    records = run_bill_agent([str(p)])

    joined = "\n".join(records[0]["audit_trail"]["ai_reasoning"])
    assert "502 Server Error" in joined
    assert "gateway.pltplt.com/ai/process" in joined
    assert "api_token" not in joined
    assert "SECRETTOKEN123" not in joined


def test_parse_error_cleaner_redacts_query_and_temp_paths():
    """parse_errors 收口：URL query 与 bill_att_* 临时路径一并脱敏（防御性）。"""
    from bill_agent.graph.nodes.parse import _clean_error

    msg = (
        "fail C:\\Users\\Jack\\AppData\\Local\\Temp\\bill_att_qv3oudgt\\data.bin "
        "https://gateway.pltplt.com/ai/process?api_token=SECRETTOKEN123&x=1"
    )
    cleaned = _clean_error(msg)
    assert "bill_att_" not in cleaned
    assert "<临时附件路径>" in cleaned
    assert "api_token" not in cleaned and "SECRETTOKEN123" not in cleaned
    assert "gateway.pltplt.com/ai/process" in cleaned


def test_eml_successful_attachment_does_not_leak_temp_path(monkeypatch, tmp_path):
    """成功解析的附件：审计轨迹用附件文件名，临时目录路径不得进入 original_text。"""
    import base64
    import io

    from docx import Document
    from tests.test_e2e_mock import _mock_extract, _result

    _mock_extract(monkeypatch, _result())

    buf = io.BytesIO()
    d = Document()
    d.add_paragraph("Handling Fee USD 100 PP")
    d.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()

    boundary = "BOUNDARY404"
    eml = tmp_path / "mail.eml"
    eml.write_bytes(
        (
            f"Content-Type: multipart/mixed; boundary={boundary}\r\n"
            "Subject: PRE-ALERT\r\nFrom: a@b.com\r\nTo: c@d.com\r\n"
            f"\r\n--{boundary}\r\nContent-Type: text/plain\r\n\r\nInvoice attached.\r\n"
            f'--{boundary}\r\nContent-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document; name="发票.docx"\r\n'
            f'Content-Disposition: attachment; filename="发票.docx"\r\nContent-Transfer-Encoding: base64\r\n\r\n{b64}\r\n'
            f"--{boundary}--\r\n"
        ).encode("utf-8")
    )

    records = run_bill_agent([str(eml)], [], "A Shipping")

    assert [r["operation"] for r in records] == ["CREATE"]
    original = records[0]["audit_trail"]["original_text"]
    # 附件 doc 的 source_file 为「父邮件名/附件名」：跨邮件同名附件才能靠它区分
    assert "=== [word] mail/发票.docx ===" in original
    _assert_no_temp_leak(records)


def test_empty_eml_without_attachments_skips_extract_and_states_cause(monkeypatch, tmp_path):
    """空正文 EML（无附件、无解析错误）：空文本不喂 LLM，转人工并写明「无内容可提取」。"""
    import bill_agent.extraction.extractor as ext

    calls = []

    def spy_llm(*a, **k):
        calls.append(a)
        raise RuntimeError("空文本不应触发 LLM")

    monkeypatch.setattr(ext, "structured_llm", spy_llm)

    # 用 write_bytes：Windows 上 write_text 会把 \r\n 翻译成 \r\r\n，破坏 MIME 分隔符
    eml = tmp_path / "empty.eml"
    eml.write_bytes(
        b"Content-Type: text/plain\r\nSubject: empty\r\nFrom: a@b.com\r\nTo: c@d.com\r\n\r\n"
    )

    records = run_bill_agent([str(eml)])

    assert calls == [], "空文本直通 pending_review，不得调用 LLM"
    assert len(records) == 1
    assert records[0]["operation"] == "PENDING_REVIEW"
    joined = "\n".join(records[0]["audit_trail"]["ai_reasoning"])
    assert "文件无内容可提取" in joined


def test_empty_body_eml_with_failed_attachment_skips_extract(monkeypatch, tmp_path):
    """EML 正文为空且附件全败：空文本不得喂 LLM（幻觉风险），直通 PENDING_REVIEW。"""
    import bill_agent.extraction.extractor as ext
    import bill_agent.parsers.ocr as ocr_mod

    def boom(path):
        raise RuntimeError("OSS 权限不足：当前 RAM 用户对桶 rag-data-2026 没有 oss:PutObject 权限")

    monkeypatch.setattr(ocr_mod, "parse_ocr", boom)
    calls = []

    def spy_llm(*a, **k):
        calls.append(a)
        raise RuntimeError("空文本不应触发 LLM")

    monkeypatch.setattr(ext, "structured_llm", spy_llm)

    boundary = "BOUNDARY202"
    eml = tmp_path / "mail.eml"
    eml.write_bytes(
        (
            f"Content-Type: multipart/mixed; boundary={boundary}\r\n"
            "Subject: PRE-ALERT\r\nFrom: a@b.com\r\nTo: c@d.com\r\n"
            f'\r\n--{boundary}\r\nContent-Type: image/png\r\nContent-Disposition: attachment; filename="发票.png"\r\n\r\n'
        ).encode("utf-8")
        + b"\x89PNG\r\n"
        + f"\r\n--{boundary}--\r\n".encode("utf-8")
    )

    records = run_bill_agent([str(eml)])

    assert calls == [], "空文本直通 pending_review，不得调用 LLM"
    assert len(records) == 1
    assert records[0]["operation"] == "PENDING_REVIEW"
    joined = "\n".join(records[0]["audit_trail"]["ai_reasoning"])
    assert "发票.png" in joined and "oss:PutObject" in joined


def test_mid_confidence_goes_through_second_confirm(monkeypatch, tmp_path):
    """中置信度（T_low ≤ c < T_high）必须经过 second_confirm 节点。"""
    import bill_agent.graph.nodes.validate_extract as ve

    calls = []

    def fake_semantic_check(result, doc_text):
        calls.append(result)
        return None

    monkeypatch.setattr(ve, "semantic_check", fake_semantic_check)

    from tests.test_e2e_mock import _mock_extract, _result, _write_docx

    _mock_extract(monkeypatch, _result(confidence=0.75))
    records = run_bill_agent([_write_docx(tmp_path)], [], "A Shipping")

    assert calls, "second_confirm 未被调用"
    assert records[0]["operation"] == "CREATE"


def test_merge_unique_keeps_distinct_content_with_same_source_and_prefix():
    """去重键是完整内容签名：同 source_file + 相同 200 字符前缀（或 excel 的空 raw_text）
    但内容不同的文档必须全部保留，否则内容会从 LLM prompt 静默消失。"""
    from bill_agent.graph.state import _merge_unique

    # 同文件多页 OCR：前 200 字符相同、尾部不同 → 两页都保留
    a = {"source_file": "x.png", "raw_text": "Z" * 250 + "A", "tables": [], "meta": {}}
    b = {"source_file": "x.png", "raw_text": "Z" * 250 + "B", "tables": [], "meta": {}}
    assert len(_merge_unique([], [a, b])) == 2

    # 同名 excel 附件：raw_text 恒为空串，表格内容不同 → 都保留
    c = {"source_file": "费用清单.xlsx", "raw_text": "", "tables": [["FEE", 100]], "meta": {}}
    d = {"source_file": "费用清单.xlsx", "raw_text": "", "tables": [["FEE", 200]], "meta": {}}
    assert len(_merge_unique([], [c, d])) == 2

    # 只有「同文件 + 完全相同内容」才是真重复
    assert len(_merge_unique([], [a, dict(a)])) == 1


def test_cross_file_same_name_attachments_both_reach_extract(monkeypatch, tmp_path):
    """两封 EML 各带同名 xlsx 附件（raw_text 均为空串）：旧去重键 (source_file, raw_text[:200])
    会把第二封的内容静默丢弃——两份费用内容都必须进入 LLM prompt。"""
    import base64
    import io

    import openpyxl

    import bill_agent.extraction.extractor as ext
    from tests.test_e2e_mock import _result

    def make_xlsx(marker: str) -> str:
        buf = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["费用名称", "金额", "币种"])
        ws.append([f"Fee {marker}", 100, "USD"])
        wb.save(buf)
        return base64.b64encode(buf.getvalue()).decode()

    prompts: list[str] = []

    class _CapturingChain:
        def __init__(self, result):
            self._result = result

        def invoke(self, messages):
            prompts.append(" ".join(getattr(m, "content", "") for m in messages))
            return self._result

    monkeypatch.setattr(ext, "structured_llm", lambda *a, **k: _CapturingChain(_result()))

    boundary = "BOUNDARY505"
    for name, marker in (("m1.eml", "ALPHA"), ("m2.eml", "BETA")):
        p = tmp_path / name
        p.write_bytes(
            (
                f"Content-Type: multipart/mixed; boundary={boundary}\r\n"
                f"Subject: Bill {marker}\r\nFrom: a@b.com\r\nTo: c@d.com\r\n"
                f"\r\n--{boundary}\r\nContent-Type: text/plain\r\n\r\nBill attached.\r\n"
                f'--{boundary}\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet; name="费用清单.xlsx"\r\n'
                f'Content-Disposition: attachment; filename="费用清单.xlsx"\r\nContent-Transfer-Encoding: base64\r\n\r\n{make_xlsx(marker)}\r\n'
                f"--{boundary}--\r\n"
            ).encode("utf-8")
        )

    records = run_bill_agent([str(tmp_path / "m1.eml"), str(tmp_path / "m2.eml")], [], "A Shipping")

    joined = "\n".join(prompts)
    assert "Fee ALPHA" in joined, "第一封 EML 的附件内容必须进入 LLM prompt"
    assert "Fee BETA" in joined, "第二封 EML 的同名附件内容不得被 fan-in 去重静默丢弃"
    assert [r["operation"] for r in records] == ["CREATE"]


def test_nested_eml_attachment_content_reaches_extract(monkeypatch, tmp_path):
    """嵌套 EML（message/rfc822 附件）：旧代码静默跳过附件、内层内容还混进外层正文，
    最终 PENDING_REVIEW 只有「未提取到任何费用项」。内层正文必须进入 LLM prompt。"""
    import bill_agent.extraction.extractor as ext
    from tests.test_e2e_mock import _result

    prompts: list[str] = []

    class _CapturingChain:
        def __init__(self, result):
            self._result = result

        def invoke(self, messages):
            prompts.append(" ".join(getattr(m, "content", "") for m in messages))
            return self._result

    monkeypatch.setattr(ext, "structured_llm", lambda *a, **k: _CapturingChain(_result()))

    inner = b"Content-Type: text/plain\r\nSubject: inner bill\r\n\r\nInner text USD 50\r\n"
    boundary = "BOUNDARY606"
    eml = tmp_path / "mail.eml"
    eml.write_bytes(
        (
            f"Content-Type: multipart/mixed; boundary={boundary}\r\n"
            "Subject: Fwd\r\nFrom: a@b.com\r\nTo: c@d.com\r\n"
            f'\r\n--{boundary}\r\nContent-Type: message/rfc822; name="fwd.eml"\r\n'
            f'Content-Disposition: attachment; filename="fwd.eml"\r\n\r\n'
        ).encode("utf-8")
        + inner
        + f"\r\n--{boundary}--\r\n".encode("utf-8")
    )

    records = run_bill_agent([str(eml)], [], "A Shipping")

    joined = "\n".join(prompts)
    assert "Inner text USD 50" in joined, "嵌套 EML 的内层正文必须进入 LLM prompt"
    # 判别点：内层邮件必须以独立文档进入 prompt（旧 walk 代码会把内层正文混进外层正文，
    # 只有外层文档的头部，永远不会有这个独立头）
    assert "=== [eml] mail/fwd.eml ===" in joined, "嵌套 EML 必须作为独立文档进入 LLM prompt"
    assert records and records[0]["operation"] == "CREATE"


def test_extraction_error_reasoning_is_redacted(monkeypatch, tmp_path):
    """提取失败信息带 URL query：进 PENDING_REVIEW 审计轨迹前必须脱敏（与 parse_errors 同收口）。"""
    import bill_agent.extraction.extractor as ext
    from tests.test_e2e_mock import _write_docx

    def boom(*a, **k):
        raise RuntimeError("LLM boom https://api.example.com/call?api_key=SECRETKEY123")

    monkeypatch.setattr(ext, "structured_llm", boom)

    records = run_bill_agent([_write_docx(tmp_path)], [], "A Shipping")

    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert pending, "提取失败重试耗尽后必须转 PENDING_REVIEW"
    joined = "\n".join(pending[0]["audit_trail"]["ai_reasoning"])
    assert "api.example.com/call" in joined, "真实原因保留"
    assert "api_key" not in joined and "SECRETKEY123" not in joined, "token 必须脱敏"
    # 重试轮次会重复累加同一消息：审计轨迹里同一异常只出现一次
    assert joined.count("LLM 提取异常") == 1


def test_doc_without_libreoffice_surfaces_real_error(monkeypatch, tmp_path):
    """.doc 无 LibreOffice 转换失败：meta.error 必须提升为 parse_errors 进审计轨迹，
    不能只剩「文件无内容可提取」的泛泛原因。"""
    import bill_agent.parsers.word as word_mod

    monkeypatch.setattr(word_mod.shutil, "which", lambda name: None)
    p = tmp_path / "旧账单.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0")  # OLE 头即可，走真实转换失败路径

    records = run_bill_agent([str(p)])

    assert len(records) == 1
    assert records[0]["operation"] == "PENDING_REVIEW"
    joined = "\n".join(records[0]["audit_trail"]["ai_reasoning"])
    assert "LibreOffice" in joined and str(p) in joined, "真实原因与文件名必须可追溯"
    assert "文件无内容可提取" not in joined


def test_doc_conversion_crash_surfaces_real_cause(monkeypatch, tmp_path):
    """LibreOffice 存在但转换崩溃/超时：不得伪装成「当前环境未安装」。"""
    import bill_agent.parsers.word as word_mod

    monkeypatch.setattr(word_mod.shutil, "which", lambda name: "soffice")

    def boom(*args, **kwargs):
        raise OSError("soffice crashed on corrupt file")

    monkeypatch.setattr(word_mod.subprocess, "run", boom)
    p = tmp_path / "旧账单.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0")

    records = run_bill_agent([str(p)])

    assert records[0]["operation"] == "PENDING_REVIEW"
    joined = "\n".join(records[0]["audit_trail"]["ai_reasoning"])
    assert "soffice crashed on corrupt file" in joined, "真实转换失败原因必须可追溯"
    assert "未安装" not in joined


def test_doc_conversion_does_not_leak_temp_path(monkeypatch, tmp_path):
    """.doc 转 docx 的落盘目录（bill_doc_）不得进入 source_file / original_text。"""
    import io
    import tempfile
    from pathlib import Path

    from docx import Document

    import bill_agent.parsers.word as word_mod
    from tests.test_e2e_mock import _mock_extract, _result

    _mock_extract(monkeypatch, _result())

    def fake_convert(path):
        buf = io.BytesIO()
        d = Document()
        d.add_paragraph("Handling Fee USD 100 PP")
        d.save(buf)
        outdir = Path(tempfile.mkdtemp(prefix="bill_doc_"))
        converted = outdir / "converted.docx"
        converted.write_bytes(buf.getvalue())
        return str(converted)

    monkeypatch.setattr(word_mod, "_convert_doc_to_docx", fake_convert)
    p = tmp_path / "账单.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0")

    records = run_bill_agent([str(p)], [], "A Shipping")

    assert records[0]["operation"] == "CREATE"
    trail = records[0]["audit_trail"]
    assert str(p) in trail["original_text"], "审计轨迹应显示原始 .doc 路径"
    blob = str(trail["source_file"]) + str(trail["original_text"])
    assert "bill_doc_" not in blob, "转换临时目录路径不得泄漏"


def test_second_confirm_llm_failure_is_not_silent_pass(monkeypatch, tmp_path):
    """语义复核调用异常不得静默当成通过（会产出「从未复核」的 CREATE），
    必须转人工并写明真实原因。"""
    import bill_agent.graph.nodes.validate_extract as ve
    from tests.test_e2e_mock import _mock_extract, _result, _write_docx

    _mock_extract(monkeypatch, _result(confidence=0.75))

    def boom(result, doc_text):
        raise RuntimeError("reviewer endpoint down")

    monkeypatch.setattr(ve, "semantic_check", boom)

    records = run_bill_agent([_write_docx(tmp_path)], [], "A Shipping")

    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert pending, "复核不可用不得静默放行"
    joined = "\n".join(pending[0]["audit_trail"]["ai_reasoning"])
    assert "reviewer endpoint down" in joined, "复核不可用的真实原因必须可追溯"


def test_final_schema_failure_not_masked_by_stale_retry_issues():
    """终态 JSON/schema 校验失败是最终原因：早先重试轮次遗留的旧问题不得冒充它。"""
    from bill_agent.graph.nodes.pending_review import pending_review_node
    from bill_agent.graph.nodes.validate_final import validate_final_node

    state = {
        "iteration": 2,
        "parsed_documents": [
            {"source_file": "mail/a.pdf", "raw_text": "费用", "tables": [], "meta": {}}
        ],
        "validation_status": "pass",
        "extraction_confidence": {"overall": 0.95},
        # 第 1 轮提取异常遗留（add 累加器保留），第 2 轮校验通过
        "validation_issues": [{"msg": "第1轮: LLM 提取异常: boom", "iteration": 1}],
        "output_records": [{"operation": "CREATE"}],  # 缺必填字段，schema 校验必失败
    }
    final = validate_final_node(state)
    assert final["validation_status"] == "fail"
    # 模拟 LangGraph add 累加语义
    state["validation_issues"] = state["validation_issues"] + final["validation_issues"]
    state["validation_status"] = final["validation_status"]

    record = pending_review_node(state)["output_records"][0]
    joined = "\n".join(record["audit_trail"]["ai_reasoning"])
    assert "输出 JSON 非法" in joined, "最终原因必须进审计轨迹"
    assert "第1轮" not in joined, "旧轮次错误不得冒充最终原因"


def test_temp_path_with_spaces_is_redacted():
    """用户名含空格时（C:\\Users\\Zhang Wei\\...）临时附件路径同样必须脱敏。"""
    from bill_agent.graph.nodes.parse import _clean_error

    msg = r"附件解析失败: C:\Users\Zhang Wei\AppData\Local\Temp\bill_att_qv3oudgt\data.bin"
    out = _clean_error(msg)
    assert "bill_att_" not in out and "Zhang Wei" not in out
    assert out.startswith("附件解析失败: <临时附件路径>")


def test_temp_path_redaction_covers_real_mkdtemp_shapes():
    """mkdtemp 随机后缀字母表含下划线（约 1/36 以 _ 开头），POSIX/UNC 路径同样必须脱敏。"""
    from bill_agent.graph.nodes.parse import _clean_error

    cases = [
        r"C:\Users\Jack\AppData\Local\Temp\bill_att__leak01\bill.docx",  # 下划线开头后缀
        r"D:\tmp\bill_doc_ab_12cd\x.docx",  # 下划线在中间
        "/tmp/bill_att_x9y8z7/data.bin",  # POSIX
        r"\\srv\share\tmp\bill_doc_z1\f.docx",  # UNC
        r"\Users\Jack\AppData\Local\Temp\bill_att_x\a.pdf",  # Windows 根相对（单反斜杠）
    ]
    for raw in cases:
        out = _clean_error(f"解析失败: {raw}")
        assert "bill_att_" not in out and "bill_doc_" not in out, out
        assert "<临时附件路径>" in out
        # 尾部路径段一并吞掉，不残留用户名/目录
        assert "Temp" not in out and "tmp" not in out and "share" not in out, out
        assert "Jack" not in out, out


def test_ocr_degradation_reason_reaches_output_reasoning():
    """OCR 降级（504 等）即使 original_text 被 500 字截断，也必须在 ai_reasoning 可追溯。"""
    from bill_agent.graph.nodes.output import output_node

    docs = [
        {
            "type": "pdf_text",
            "source_file": "a.pdf",
            "raw_text": "X" * 600,
            "tables": [],
            "meta": {"ocr_error": "OCR 返回失败: 504 Gateway Timeout", "ocr_fallback": "pdf_text"},
        }
    ]
    state = {
        "parsed_documents": docs,
        "match_results": [
            {"operation": "CREATE", "charge": {"charge_item": "OCEAN FREIGHT", "original_charge_item": "O/F"}}
        ],
        "posting_type": "AR",
        "ar_ap_reasoning": "AR",
        "accounting_entity": "A Shipping",
    }
    record = output_node(state)["output_records"][0]
    joined = "\n".join(record["audit_trail"]["ai_reasoning"])
    assert "504" in joined and "降级" in joined, "降级原因必须进审计轨迹，不能因截断消失"


def test_pure_inline_cid_email_reasoning_mentions_inline_images():
    """正文为空、只有 cid 内联图的邮件：PENDING_REVIEW 必须说明图片为何未解析。"""
    from bill_agent.graph.nodes.pending_review import pending_review_node

    state = {
        "parsed_documents": [
            {
                "type": "eml",
                "source_file": "mail.eml",
                "raw_text": "",
                "tables": [],
                "meta": {
                    "inline_cid_parts": [
                        {"filename": "InsertPic_A1.jpg", "content_type": "image/jpeg"}
                    ]
                },
            }
        ],
        "validation_status": "pass",
        "extraction_confidence": {"overall": 0.0},
    }
    record = pending_review_node(state)["output_records"][0]
    joined = "\n".join(record["audit_trail"]["ai_reasoning"])
    assert "InsertPic_A1.jpg" in joined, "内联图未被解析的原因必须可追溯"
    assert "无内容可提取" in joined


def test_semantic_check_empty_response_is_issue(monkeypatch):
    """模型返回空/空白内容不等于复核通过。"""
    from bill_agent.validation import llm_validator

    class _Resp:
        content = "   "

    class _LLM:
        def invoke(self, msg):
            return _Resp()

    monkeypatch.setattr(llm_validator, "get_llm", lambda: _LLM())

    issue = llm_validator.semantic_check({"charges": []}, "原文")
    assert issue and "未返回" in issue


def test_second_confirm_empty_review_goes_pending(monkeypatch, tmp_path):
    """语义复核返回空响应：中置信度单据必须转人工，不能静默 CREATE。"""
    from bill_agent.validation import llm_validator
    from tests.test_e2e_mock import _mock_extract, _result, _write_docx

    _mock_extract(monkeypatch, _result(confidence=0.75))

    class _Resp:
        content = ""

    class _LLM:
        def invoke(self, msg):
            return _Resp()

    monkeypatch.setattr(llm_validator, "get_llm", lambda: _LLM())

    records = run_bill_agent([_write_docx(tmp_path)], [], "A Shipping")

    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert pending, "复核未返回内容不得静默放行"
    joined = "\n".join(pending[0]["audit_trail"]["ai_reasoning"])
    assert "未返回" in joined


def test_second_confirm_failure_reason_reaches_audit_trail(monkeypatch, tmp_path):
    """语义复核发现的问题必须进 PENDING_REVIEW 审计轨迹，不能只剩泛泛的「低置信度或校验失败」。"""
    import bill_agent.graph.nodes.validate_extract as ve
    from tests.test_e2e_mock import _mock_extract, _result, _write_docx

    _mock_extract(monkeypatch, _result(confidence=0.75))
    monkeypatch.setattr(ve, "semantic_check", lambda result, doc_text: "总金额与明细不符，疑似虚构费用")

    records = run_bill_agent([_write_docx(tmp_path)], [], "A Shipping")

    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert pending, "语义复核失败必须转人工"
    joined = "\n".join(pending[0]["audit_trail"]["ai_reasoning"])
    assert "总金额与明细不符" in joined, "复核发现的问题必须可追溯"


def test_low_confidence_after_retry_shows_confidence_not_stale_errors(monkeypatch, tmp_path):
    """重试后校验通过但置信度不足：审计轨迹写明真实原因（置信度数值），
    历史轮次的旧错误不得冒充最终原因。"""
    import bill_agent.extraction.extractor as ext
    from tests.test_e2e_mock import _result, _write_docx

    calls = []

    class _Chain:
        def __init__(self, result):
            self._result = result

        def invoke(self, messages):
            return self._result

    def flaky_llm(*a, **k):
        calls.append(1)
        if len(calls) == 1:
            raise RuntimeError("round1 boom")
        return _Chain(_result(confidence=0.5))

    monkeypatch.setattr(ext, "structured_llm", flaky_llm)

    records = run_bill_agent([_write_docx(tmp_path)], [], "A Shipping")

    pending = [r for r in records if r["operation"] == "PENDING_REVIEW"]
    assert pending
    joined = "\n".join(pending[0]["audit_trail"]["ai_reasoning"])
    assert "置信度" in joined, "必须写明置信度不足的真实原因"
    assert "round1 boom" not in joined, "历史轮次的旧错误不得冒充最终原因"


def test_empty_xlsx_skips_extract(monkeypatch, tmp_path):
    """空工作簿（tables=[[]]）不算内容：空文本不得喂 LLM，直通 PENDING_REVIEW。"""
    import openpyxl

    import bill_agent.extraction.extractor as ext

    calls = []

    def spy_llm(*a, **k):
        calls.append(a)
        raise RuntimeError("空文本不应触发 LLM")

    monkeypatch.setattr(ext, "structured_llm", spy_llm)

    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.save(p)

    records = run_bill_agent([str(p)])

    assert calls == [], "空工作簿直通 pending_review，不得调用 LLM"
    assert records[0]["operation"] == "PENDING_REVIEW"
    assert "文件无内容可提取" in "\n".join(records[0]["audit_trail"]["ai_reasoning"])


def test_temp_path_redaction_is_conservative_on_multi_path_lines():
    """同一行出现 marker 时，从行内最早路径起点整段脱敏（宁可过脱敏）。

    旧实现尝试在「语句边界」处切开以保留前一句的普通路径，但边界判定会被
    路径自身骗过——组件含空格后紧跟分隔符（'CI, Build Agent \\bill_att_x'）时，
    该分隔符看起来就是新语句起点，用户名整段留在明文（round 13 blocker）。
    现在一律取行内最早起点：前一句路径也一并隐去，换取任何形态都不泄漏。
    """
    from bill_agent.graph.nodes.parse import _clean_error

    msg = r"日志 C:\Windows\notepad.exe；附件 C:\Temp\bill_att_x\a.pdf"
    out = _clean_error(msg)

    assert "bill_att_" not in out and "Temp" not in out
    assert "notepad.exe" not in out, "同一行更早的路径也一并隐去（保守策略）"
    assert out == "日志 <临时附件路径>"


def test_temp_path_redaction_stops_exactly_at_marker_end():
    """替换区间必须**排除**紧跟 marker 的空格/引号：否则会把下一个词也删掉。

    日志里路径后面通常直接跟空白（'...bill_att_x failed to read'）或引号；
    终点条件写成「小于等于」会把 'failed' 或收尾引号一并吞掉，审计文本失真。
    """
    from bill_agent.graph.nodes.parse import _clean_error

    out = _clean_error("open C:\\Temp\\bill_att_x failed to read")
    assert out == "open <临时附件路径> failed to read", out

    out = _clean_error("Failed to open 'C:\\Temp\\bill_att_x' then retry")
    assert out == "Failed to open '<临时附件路径>' then retry", out


def test_temp_path_redaction_rejects_glued_root_anchors():
    """根被前一个词粘住时（'见/home/jack/...'、'see/home/...'），路径内部的
    斜杠不能当起点——否则用户名/目录前缀留在明文（round 14）。

    CJK 粘连的根能被正则认出（非 ASCII 左邻不排除），但左邻不是外部边界；
    ASCII 粘连的根根本认不出。两者都必须退化为整行脱敏。路径组件以可信标点
    结尾（'foo('、'ci:'）时同理：起点之前已有分隔符，那个斜杠只是路径内部的
    分隔符。
    """
    from bill_agent.graph.nodes.parse import _clean_error

    cases = [
        ("见/home/jack/Program Files (x86)/bill_att_x/a.pdf", ("jack", "Program", "home")),
        ("见/tmp/K2 /bill_att_x/a.pdf", ("K2", "tmp")),
        ("见/home/John Smith/tmp/bill_att_x/a.pdf", ("John", "Smith", "home")),
        ("见/home/jack/foo(/bill_att_x/a.pdf", ("jack", "foo", "home")),
        ("见/home/jack/ci:/bill_att_x/a.pdf", ("jack", "ci", "home")),
        ("see/home/jack/Program Files (x86)/bill_att_x/a.pdf", ("jack", "Program", "home")),
        ("see/tmp/K2 /bill_att_x/a.pdf", ("K2", "tmp")),
        # 'C:\\' 被 '见' 挡住后，盘符冒号仍是可信边界：最多留下 'C:' 本身，
        # Users/Jack 等必须隐去
        (r"见C:\Users\Jack\AppData\Local\Temp\bill_att_x\a.pdf",
         ("Jack", "Users", "AppData", "Temp")),
        # 行内更早的普通 token 带分隔符（N/A）：没有可信起点，整行脱敏
        (r"状态 N/A，附件 C:\Temp\bill_att_x\a.pdf", ("N/A", "Temp")),
    ]
    for msg, sensitive in cases:
        out = _clean_error(msg)
        assert "<临时附件路径>" in out, (msg, out)
        assert not any(t in out for t in sensitive), (msg, out)


def test_temp_path_redaction_handles_boundary_chars_inside_path():
    """路径前缀本身含 , ; :（Smith, John、srv,1、ci:run）：这些不是语句边界，
    必须从盘符/UNC/POSIX 起点整段脱敏，不能把用户名或共享名留在明文。"""
    from bill_agent.graph.nodes.parse import _clean_error

    cases = [
        r"解析失败: C:\Users\Smith, John\AppData\Local\Temp\bill_att_x\a.pdf",
        r"解析失败: \\srv,1\share\Temp\bill_att_x\a.pdf",
        "解析失败: /var/tmp/ci:run/bill_att_x/a.pdf",
        # 逗号后是「无分隔符的续词」再才是路径延续：旧实现把续词当外部语句，
        # 从 marker 起算，用户名/自定义 TEMP 目录整段留在明文
        r"C:\Users\Smith, John Doe\AppData\Local\Temp\bill_att_x\a.pdf",
        r"D:\CI, Build Agent\tmp\bill_att_x\a.pdf",
        "/home/Smith, John Doe/tmp/bill_att_x/a.pdf",
        r"C:\Users\Smith;John Doe\AppData\Local\Temp\bill_att_x\a.pdf",
    ]
    for msg in cases:
        out = _clean_error(msg)
        assert "<临时附件路径>" in out, out
        assert "bill_att_" not in out, out
        assert not any(
            s in out
            for s in (
                "Smith", "srv", "share", "ci:run", "Temp", "tmp",
                "John", "Doe", "AppData", "Users", "CI", "Build Agent", "home",
            )
        ), out


def test_temp_path_redaction_anchors_paths_after_punctuation():
    """POSIX/根相对路径常紧跟 ：、=、: 且无空格（失败：/home/...、TMPDIR=/...、
    failed:/...）：这些也是路径起点，不能从 marker 起算把 /home/用户名 留在明文。"""
    from bill_agent.graph.nodes.parse import _clean_error

    cases = [
        "失败：/home/alice/tmp/bill_att_x/a.pdf",
        "TMPDIR=/home/alice/tmp/bill_att_x/a.pdf",
        "failed:/home/alice/tmp/bill_att_x/a.pdf",
        r"failed:\Users\Jack\AppData\Local\Temp\bill_att_x\a.pdf",
    ]
    for msg in cases:
        out = _clean_error(msg)
        assert "<临时附件路径>" in out, out
        assert not any(
            s in out for s in ("alice", "home", "Users", "Jack", "AppData", "Temp", "bill_att_")
        ), out


def test_temp_path_redaction_marker_is_case_insensitive():
    """Windows 工具链会把路径转成大写：BILL_ATT_* 同样是临时落盘路径。"""
    from bill_agent.graph.nodes.parse import _clean_error

    out = _clean_error(r"C:\Users\Jack\AppData\Local\Temp\BILL_ATT_X\a.pdf")
    assert "<临时附件路径>" in out, out
    assert not any(s in out for s in ("Jack", "AppData", "Temp", "BILL_ATT_")), out


def test_parse_error_sink_redacts_all_credential_url_forms():
    """parse_errors 是真实收口：协议相对 userinfo 与无 scheme fragment 都必须清掉。"""
    from bill_agent.graph.nodes.parse import _error

    state = {"source_file": {"path": "a.pdf"}}
    msg = _error(state, RuntimeError("boom //sinkuser:SINKPASS_ABC123@gateway.invalid/p?x=1"))
    err = msg["parse_errors"][0]["error"]
    assert "SINKPASS_ABC123" not in err and "sinkuser" not in err, err

    msg = _error(state, RuntimeError("POST /ai/process#api_token=FRAGTOKEN987 failed"))
    err = msg["parse_errors"][0]["error"]
    assert "FRAGTOKEN987" not in err, err
    assert "/ai/process" in err, "保留路径便于定位故障"

    # fragment 在 query 之前：只看 '?' 会把 fragment 里的凭据留下（round 13）
    msg = _error(state, RuntimeError("POST /ai/process#api_token=FRAGTOKEN987?x=1 failed"))
    err = msg["parse_errors"][0]["error"]
    assert "FRAGTOKEN987" not in err and "api_token" not in err, err
    assert "/ai/process" in err


def test_temp_path_redaction_covers_repr_escaped_paths():
    """pymupdf 等库输出的路径是 repr/JSON 转义形式（分隔符成对）：
    必须从盘符/UNC 起点整段脱敏，不能把用户名与临时目录留在明文。"""
    import json

    from bill_agent.graph.nodes.parse import _clean_error

    msg = "Failed to open file " + repr(r"C:\Users\Jack\AppData\Local\Temp\bill_att_ab12cd\invoice.pdf")
    out = _clean_error(msg)
    assert "<临时附件路径>" in out
    assert not any(s in out for s in ("Users", "Jack", "AppData", "Temp", "bill_att_")), out

    unc = '{"path": ' + json.dumps(r"\\srv\share\tmp\bill_doc_z1\f.docx") + "}"
    out = _clean_error(unc)
    assert "<临时附件路径>" in out
    assert not any(s in out for s in ("srv", "share", "bill_doc_")), out

    # round 13：转义路径里组件含逗号 + 空格，边界后的双反斜杠曾被当成新起点，
    # '...CI, Build Agent' 整段留在明文
    escaped = "Failed to open " + repr(
        r"C:\Users\Jack\AppData\Local\Temp\audit_r13\CI, Build Agent\bill_att_ab12cd34\invoice.pdf"
    )
    out = _clean_error(escaped)
    assert "<临时附件路径>" in out
    assert not any(s in out for s in ("Jack", "audit_r13", "CI", "Build Agent", "Temp", "bill_att_")), out

    escaped_cjk = "OSError('" + r"C:\Users\张三,李四\AppData\Local\Temp\bill_att_x\a.pdf".replace(
        "\\", "\\\\"
    ) + "')"
    out = _clean_error(escaped_cjk)
    assert "<临时附件路径>" in out
    assert not any(s in out for s in ("张三", "李四", "AppData", "Temp", "bill_att_")), out


def test_temp_path_redaction_covers_spaced_and_unicode_components():
    """组件含空格/分号/CJK/括号、根相对路径无盘符：都必须从组件开头整段脱敏。"""
    from bill_agent.graph.nodes.parse import _clean_error

    messages = [
        "解析失败: C:\\Users\\Smith,\\AppData\\Local\\Temp\\bill_att_abcd1234\\a.pdf",
        r"failed:\Users\Zhang Wei\AppData\Local\Temp\bill_att_abcd1234\a.pdf",
        "解析失败: C:\\Users\\张三,李四\\AppData\\Local\\Temp\\bill_att_x\\a.pdf",
        r"C:\Users\Smith;John Doe\AppData\Local\Temp\bill_att_x\a.pdf",
        r"C:\Users\Bob (old, 2020)\AppData\Local\Temp\bill_att_x\a.pdf",
        r"D:\CI, Build Agent\tmp\bill_att_x\a.pdf",
        "失败：/home/张三 李四/tmp/bill_att_x/a.pdf",
    ]
    for msg in messages:
        out = _clean_error(msg)
        assert "<临时附件路径>" in out, out
        assert not any(
            s in out
            for s in (
                "Smith", "Zhang", "Wei", "张三", "李四", "Bob", "old", "CI",
                "Build Agent", "AppData", "Temp", "home", "bill_att_",
            )
        ), out


def test_temp_path_redaction_is_linear_on_separator_heavy_input():
    """分隔符密集输入不得让脱敏退化为二次方。

    病理输入分别覆盖四代回归：
    - 32k '/' + 换行后的路径：旧惰性前缀正则需 4.6s（marker 不可达时逐起点回溯）；
    - 64k '\\' + 同一行的 marker：repr 转义路径里每个双反斜杠都被 UNC 分支命中，
      取最后一个起点会逐 match 搜 gap，实测二次方；
    - 32k 语句边界符（':', '；'）同一行：旧实现为每条边界向右扫到空白判 token，
      64k 个 ':' 实测 94s；
    - 64k 个字母（无 ://）与 '/+' 重复串：URL 正则若对每个起点向后找 :// 或 ?/#，
      同样退化为二次方；
    - marker 密集且无空白分隔：旧实现每个 marker 都用正则向后扫到行尾（找引号/
      空格的终点），200KB 实测 5.5s、400KB 22s。
    """
    import gc
    import time

    from bill_agent.graph.nodes.parse import _clean_error

    def _run(text, reps: int = 3):
        """取多轮最小值：GC/调度噪声会让单次计时在毫秒级抖动两个数量级。"""
        best = float("inf")
        out = ""
        for _ in range(reps):
            gc.collect()
            started = time.perf_counter()
            out = _clean_error(text)
            best = min(best, time.perf_counter() - started)
        return best, out

    cases = [
        "/" * 32_000 + "\n/tmp/bill_att_x/a.pdf",
        "\\" * 64_000 + " bill_doc_x\\a.pdf",
        ":" * 32_000 + " bill_att_x/a.pdf",
        "；" * 32_000 + " bill_att_x/a.pdf",
        "a" * 64_000,
        "/+" * 32_000,
        "bill_att_a/" * 6_400,
        ("bill_att_a/" * 20_000)[:200_000],
        ("bill_att_x," * 11_000)[:110_000],
        ("；" * 100_000 + " bill_att_x/a.pdf")[:200_000],
        # marker 与终点**同时**密集：终点指针若每个 marker 都从头扫一遍，
        # 16k 个 marker+空格实测 5.5s（此前用例只有 marker 密集或终点密集）
        "bill_att_x " * 16_000,
    ]
    for text in cases:
        elapsed, out = _run(text)

        assert "bill_att_" not in out and "bill_doc_" not in out
        assert elapsed < 1.0, f"脱敏耗时 {elapsed:.2f}s，疑似回溯退化"

    # 绝对墙钟阈值在更快的机器上会失效：再钉住规模增长率（16k→64k 应≈4x）
    small_t, _ = _run("\\" * 16_000 + " bill_doc_x\\a.pdf")
    large_t, _ = _run("\\" * 64_000 + " bill_doc_x\\a.pdf")
    assert large_t < max(8 * small_t, 0.05), (
        f"16k={small_t:.4f}s 64k={large_t:.4f}s，疑似超线性"
    )


def test_temp_path_redaction_never_leaks_generated_path_components():
    """生成式回归（round 13）：根 × 组件 × 包装组合，任何组件都不得留在明文。

    手写用例覆盖不到「组件含空格 + 转义分隔符」「组件末尾空格 + 路径延续」
    这类组合；这里用固定组合矩阵复现 round 13 实测的 7,203/20,000 泄漏形态，
    并钉住「根起点必须被认出，否则整行脱敏」的下界。
    """
    import json

    from bill_agent.graph.nodes.parse import _clean_error

    roots = [
        ("C:\\", ()),
        ("D:/", ()),
        ("\\\\srv\\share\\", ("srv", "share")),
        ("\\Users\\", ("Users",)),
        ("/home/", ("home",)),
        ("/var/tmp/", ("var", "tmp")),
    ]
    comps = [
        "Smith, John Doe",
        "CI, Build Agent",
        "srv,1",
        "ci:run",
        "张三 李四",
        "Bob (old, 2020)",
        "K2 ",
    ]
    wrappers = [
        "%s",
        "failed:%s",
        "失败：%s",
        "TMPDIR=%s",
        "[Errno 13] Permission denied: '%s'",
        "Failed to open file %s",
        '{"path": %s}',
        # 根被前一个词粘住（CJK / ASCII）：路径内部的斜杠不得成为起点
        "见%s",
        "see%s",
    ]
    checked = 0
    for root, root_tokens in roots:
        sep = "/" if "/" in root else "\\"
        for a in comps:
            for b in comps:
                path = f"{root}{a}{sep}{b}{sep}bill_att_x{sep}a.pdf"
                sensitive = root_tokens + tuple(
                    p for c in (a, b) for p in [c, *c.split()] if len(p) > 2
                )
                for raw in (
                    path,
                    repr(path),
                    json.dumps(path),
                    json.dumps(path, ensure_ascii=False),
                ):
                    for wrapper in wrappers:
                        msg = wrapper % raw
                        out = _clean_error(msg)
                        assert "<临时附件路径>" in out, msg
                        assert not any(t in out for t in sensitive), (msg, out)
                        checked += 1
    assert checked >= 500


def test_large_document_budget_estimate_uses_truncated_prompt(monkeypatch):
    """500k 字文档：真正的 prompt 截断到 30k 字，预算估算也必须按截断后计算，
    否则零次 LLM 调用就以「预算不足」转 PENDING_REVIEW。"""
    import bill_agent.extraction.extractor as ext
    from bill_agent.harness.budget import BudgetTracker
    from tests.test_e2e_mock import _result

    captured: dict[str, int] = {}

    class _CapturingChain:
        def __init__(self, result):
            self._result = result

        def invoke(self, messages):
            captured["user_len"] = len(messages[1].content)
            return self._result

    monkeypatch.setattr(ext, "structured_llm", lambda *a, **k: _CapturingChain(_result()))
    # 默认上限 200k token；500k 字全文估算 250k 会误报，截断后约 15k
    monkeypatch.setattr(ext, "get_budget", lambda: BudgetTracker(200_000, 20))

    docs = [
        {
            "type": "text",
            "source_file": "huge.txt",
            "raw_text": "A" * 500_000,
            "tables": [],
            "meta": {},
        }
    ]

    result = ext.extract(docs)

    assert result.charges, "大文档必须真正调用 LLM，而不是被预算闸门误拦"
    assert captured["user_len"] <= 30_100, "实际发送的 prompt 必须被截断到 30k 字以内"
